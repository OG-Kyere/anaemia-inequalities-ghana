"""Assemble the requested two-column reading version with every display in main."""
from copy import deepcopy
from pathlib import Path
import re
from docx import Document
from docx.shared import Inches,Pt
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

def format_main(final, root):
    path=final/'Main_Manuscript.docx'
    doc=Document(path)
    blocks=list(doc.element.body)[:-1]
    def text(el):
        return Paragraph(el,doc._body).text if el.tag==qn('w:p') else ''
    def index(label): return next(i for i,b in enumerate(blocks) if text(b)==label)
    def clean(seq):
        return [deepcopy(b) for b in seq if not (b.tag==qn('w:p') and not text(b).strip() and b.xpath('.//w:br[@w:type="page"]'))]
    intro,results,discussion,refs=[index(x) for x in ('Introduction','Results','Discussion','References')]
    t1,t2,t3,legend=[index(x) for x in ('Table I','Table II','Table III','Figure legend')]
    front=clean(blocks[:intro]); methods=clean(blocks[intro:results])
    tail=clean(blocks[discussion:refs]); references=clean(blocks[refs:t1])
    table1=clean(blocks[t1:t2]); table2=clean(blocks[t2:t3]); table3=clean(blocks[t3:legend])
    heads=['Study population and prevalence','Wealth-related inequality','Adjusted associations with anaemia','Decomposition of inequality','Community-level heterogeneity']
    positions={h:next(i for i in range(results,discussion) if text(blocks[i])==h) for h in heads}
    chunks={}
    for h in heads:
        start=positions[h]
        end=min([v for v in positions.values() if v>start]+[discussion])
        chunks[h]=clean(blocks[start:end])
    for el in list(doc.element.body):
        if el.tag!=qn('w:sectPr'): doc.element.body.remove(el)
    normal=doc.styles['Normal']
    normal.font.size=Pt(10.5); normal.paragraph_format.line_spacing=1.1
    normal.paragraph_format.space_after=Pt(4)
    for name,size in [('Title',14),('Heading 1',12),('Heading 2',11),('Heading 3',10.5)]:
        doc.styles[name].font.size=Pt(size)
        doc.styles[name].paragraph_format.space_before=Pt(7)
        doc.styles[name].paragraph_format.space_after=Pt(4)
    def columns(n):
        s=doc.sections[-1]
        s.page_width=Inches(8.2677); s.page_height=Inches(11.6929)
        s.left_margin=s.right_margin=Inches(.7)
        s.top_margin=s.bottom_margin=Inches(.7)
        c=s._sectPr.find(qn('w:cols'))
        if c is None: c=OxmlElement('w:cols'); s._sectPr.append(c)
        c.set(qn('w:num'),str(n)); c.set(qn('w:space'),'360')
        c.set(qn('w:equalWidth'),'1')
    def change(n):
        doc.add_section(WD_SECTION_START.CONTINUOUS); columns(n)
    def add(seq):
        for el in seq: doc.element.body.insert(-1,el)
    def display(seq):
        change(1); add(seq); change(2)
    def figure(n,stem,title,caption):
        change(1)
        p=doc.add_paragraph('Figure '+str(n)+'. '+title,'Heading 2')
        p.paragraph_format.keep_with_next=True
        p=doc.add_paragraph(); p.paragraph_format.keep_with_next=True
        p.add_run().add_picture(str(final/'figures'/stem),width=Inches(6.15))
        p=doc.add_paragraph(caption); p.paragraph_format.space_after=Pt(7)
        change(2)
    columns(1); add(front); change(2); add(methods)
    doc.add_paragraph('Results','Heading 1')
    add(chunks[heads[0]]); display(table1)
    figure(1,'Figure_S2_wealth_gradient.png','Wealth gradient in anaemia prevalence',
           'Survey-weighted prevalence with 95% confidence intervals across household wealth quintiles. The poorest and richest quintiles had prevalence of 46.6% and 38.8%, respectively.')
    add(chunks[heads[1]])
    figure(2,'Figure_S1_concentration_curve.png','Concentration curve for anaemia',
           'Women are ranked from poorest to richest. The dashed diagonal marks equality. Erreygers index = -0.0589; 95% stratified PSU bootstrap interval -0.0924 to -0.0249 (5,000 replicates).')
    add(chunks[heads[2]]); display(table2)
    figure(3,'Figure_1_regional_prevalence.png','Anaemia prevalence by region',
           'Survey-weighted prevalence with 95% confidence intervals in Ghana\'s 16 regions. The dashed line marks national prevalence of 41.1%; estimates range from about 30.1% in Bono to 51.8% in Oti.')
    add(chunks[heads[3]]); add(chunks[heads[4]]); display(table3)
    change(1)
    doc.add_paragraph('Table IV. Overall Wald tests in the adjusted survey-weighted model','Heading 2')
    source=(root/'results/table4_overall_wald_tests.md').read_text(encoding='utf8')
    values=[[x.strip() for x in line.strip().strip('|').split('|')] for line in source.splitlines() if line.startswith('|')]
    values=[values[0]]+values[2:]
    table=doc.add_table(rows=len(values),cols=4); table.style='Table Grid'
    for i,row in enumerate(values):
        for j,value in enumerate(row): table.cell(i,j).text=value
    # Reuse the standard editable-table formatting without rebuilding other tables.
    from docx.enum.table import WD_TABLE_ALIGNMENT
    table.alignment=WD_TABLE_ALIGNMENT.CENTER; table.autofit=False
    for j,width in enumerate([2.0,1.3,.8,2.45]):
        table.columns[j].width=Inches(width)
        for i,row in enumerate(table.rows):
            cell=row.cells[j]; cell.text=cell.text.replace('**',''); cell.width=Inches(width)
            for p in cell.paragraphs:
                p.paragraph_format.line_spacing=1.1
                p.paragraph_format.space_after=Pt(2)
                for r in p.runs: r.font.size=Pt(9.5); r.bold=(i==0)
    hdr=OxmlElement('w:tblHeader'); table.rows[0]._tr.get_or_add_trPr().append(hdr)
    for cell in table.rows[0].cells:
        sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),'EDEDED'); cell._tc.get_or_add_tcPr().append(sh)
    doc.add_paragraph('Analytic sample: 7,550 women. Tests use the survey-weighted full model with stratification and PSU clustering. df denotes degrees of freedom.')
    change(2)
    figure(4,'Figure_S3_decomposition.png','Decomposition of wealth-related anaemia inequality',
           'Absolute contributions to the complete-case Erreygers index (-0.058853; n = 7,550). Negative contributions reinforce concentration among poorer women; positive contributions offset it. Table III gives bootstrap intervals. Contributions are statistical, not causal.')
    add(tail); add(references)
    for p in doc.paragraphs:
        if p._p.xpath('.//m:oMathPara'):
            tabs=p.paragraph_format.tab_stops; tabs.clear_all()
            from docx.enum.text import WD_TAB_ALIGNMENT
            tabs.add_tab_stop(Inches(3.12),WD_TAB_ALIGNMENT.RIGHT)
        if p.text=='References': p.paragraph_format.page_break_before=True
        if re.match(r'^\d+\. ',p.text):
            p.paragraph_format.line_spacing=1.05
            for r in p.runs: r.font.size=Pt(9)
    doc.save(path)
    # Filenames now match the four main-paper figure numbers.
    names={'Figure_S2_wealth_gradient':'Figure_1_wealth_gradient','Figure_S1_concentration_curve':'Figure_2_concentration_curve','Figure_1_regional_prevalence':'Figure_3_regional_prevalence','Figure_S3_decomposition':'Figure_4_decomposition'}
    for old,new in names.items():
        for suffix in ('.png','.tiff'): (final/'figures'/(old+suffix)).rename(final/'figures'/(new+suffix))
    assert len(Document(path).tables)==5
    assert sum(1 for s in Document(path).sections if s._sectPr.find(qn('w:cols')).get(qn('w:num'))=='2')>=1

