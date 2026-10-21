from __future__ import annotations

import json
import re
import subprocess
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image

ROOT = Path(__file__).resolve().parents[2]
BUILD = ROOT/'work/review_package'
FINAL = BUILD/'final'
TEMP = BUILD/'source'
FIGS = FINAL/'figures'
for folder in (FINAL, TEMP, FIGS):
    folder.mkdir(parents=True, exist_ok=True)

def read(name):
    return (ROOT/name).read_text(encoding='utf8')

TITLE = 'Socioeconomic and Geographic Inequalities in Anaemia Among Women of Reproductive Age in Ghana'
NAMES = 'Gideon Ofosu Kyere, Wilhemina Adoma Pels, Prince Apaah and Clement Acheampong'
AFFILIATION = 'Department of Statistics and Actuarial Science, Kwame Nkrumah University of Science and Technology, Kumasi, Ghana'

def set_styles(doc):
    section=doc.sections[0]
    section.page_width=Inches(8.2677); section.page_height=Inches(11.6929)
    section.top_margin=section.bottom_margin=Inches(.85)
    section.left_margin=section.right_margin=Inches(.85)
    for style in doc.styles:
        for border in style.element.xpath('.//w:pBdr'):
            border.getparent().remove(border)
        if style.type == 1:
            style.font.name='Times New Roman'
            style.font.color.rgb=RGBColor(0,0,0)
    normal=doc.styles['Normal']
    normal.font.size=Pt(11)
    normal.paragraph_format.line_spacing=1.5
    normal.paragraph_format.space_after=Pt(6)
    for name,size in [('Title',16),('Heading 1',13),('Heading 2',12),('Heading 3',11)]:
        doc.styles[name].font.size=Pt(size)
        doc.styles[name].font.bold=True
        doc.styles[name].paragraph_format.space_before=Pt(10)
        doc.styles[name].paragraph_format.space_after=Pt(5)
    doc.core_properties.author='Gideon Ofosu Kyere and coauthors'
    doc.core_properties.title=TITLE
    doc.core_properties.subject='Anaemia inequalities in Ghana'
    doc.core_properties.comments=''

ref=Document(); set_styles(ref); ref.save(TEMP/'reference.docx')

def block_table(text,index=0):
    groups=re.findall(r'(?m)(?:^\|.*\|\s*\n)+',text)
    return groups[index].strip()+'\n'

def equation_layout(text):
    # Equivalent line wrapping for Word; equation definitions are unchanged.
    def convert(match):
        math=match.group(1).strip()
        number=re.search(r'\\tag\{(\d+)\}',math)
        tag=number.group(1) if number else None
        math=re.sub(r'\\tag\{\d+\}','',math).strip()
        math=math.replace('v005',r'\text{v005}')
        if tag in ('3','6','7','8'):
            parts=re.split(r'\\qquad\s*',math)
            math='\\begin{gathered}\n'+' \\\\\n'.join(p.strip() for p in parts)+'\n\\end{gathered}'
        return '$$\n'+math+'\n$$\n\n'+('EQUATIONNUMBER'+tag+'\n' if tag else '')
    return re.sub(r'\$\$([\s\S]*?)\$\$',convert,text)

def style_tables(doc):
    # Rebuild pandoc tables with standard Word cells before applying layout.
    # This avoids custom reference-style dependencies in headless importers.
    originals=list(doc.tables)
    rebuilt=[]
    for old in originals:
        values=[[cell.text for cell in row.cells] for row in old.rows]
        table=doc.add_table(rows=len(values),cols=len(values[0]))
        table.style='Table Grid'
        for i,row in enumerate(values):
            for j,value in enumerate(row): table.cell(i,j).text=value
        old._tbl.addnext(table._tbl)
        old._tbl.getparent().remove(old._tbl)
        rebuilt.append(table)
    for table in rebuilt:
        table.alignment=WD_TABLE_ALIGNMENT.CENTER
        table.autofit=False
        count=len(table.columns)
        widths={5:[1.25,1.6,.8,.7,2.2],4:[2.0,1.0,1.0,2.55],3:[3.35,1.6,1.6]}.get(count,[6.55/count]*count)
        first=table.cell(0,0).text
        if first=='Factor' and count==5: widths=[1.2,2.0,.85,1.35,1.15]
        for cell in table.rows[0].cells:
            if cell.text=='Unweighted n': cell.text='Unweighted\nn'
            if cell.text=='Adjusted OR': cell.text='Adjusted\nOR'
        if first=='Item': widths=[.45,1.05,2.0,3.1]
        if first=='Author' and count==4: widths=[1.2,1.6,1.65,2.1]
        for j,col in enumerate(table.columns):
            col.width=Inches(widths[j])
        borders=OxmlElement('w:tblBorders')
        for edge in ['top','left','bottom','right','insideH','insideV']:
            b=OxmlElement('w:'+edge); b.set(qn('w:val'),'single'); b.set(qn('w:sz'),'4'); b.set(qn('w:color'),'D9D9D9'); borders.append(b)
        table._tbl.tblPr.insert_element_before(borders,'w:shd','w:tblLayout','w:tblCellMar','w:tblLook')
        for i,row in enumerate(table.rows):
            no_split=OxmlElement('w:cantSplit'); row._tr.get_or_add_trPr().append(no_split)
            if i==0:
                repeat=OxmlElement('w:tblHeader'); row._tr.get_or_add_trPr().append(repeat)
            for j,cell in enumerate(row.cells):
                cell.width=Inches(widths[j]); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                margins=OxmlElement('w:tcMar')
                for side in ('top','left','bottom','right'):
                    el=OxmlElement('w:'+side); el.set(qn('w:w'),'65'); el.set(qn('w:type'),'dxa'); margins.append(el)
                cell._tc.get_or_add_tcPr().insert_element_before(margins,'w:vAlign')
                if i==0:
                    shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'EDEDED'); cell._tc.get_or_add_tcPr().append(shade)
                for p in cell.paragraphs:
                    p.paragraph_format.line_spacing=1.15
                    p.paragraph_format.space_after=Pt(1)
                    p.paragraph_format.space_before=Pt(1)
                    p.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 or (count==5 and j==1) else WD_ALIGN_PARAGRAPH.CENTER
                    for run in p.runs:
                        run.font.name='Times New Roman'; run.font.size=Pt(9.5)
                        if i==0: run.bold=True

def make_doc(name, markdown):
    source=TEMP/(name+'.md'); target=FINAL/(name+'.docx')
    source.write_text(equation_layout(markdown),encoding='utf8')
    subprocess.run(['pandoc',str(source),'--from=markdown+tex_math_dollars','--to=docx','--reference-doc='+str(TEMP/'reference.docx'),'--output='+str(target)],check=True)
    doc=Document(target); set_styles(doc)
    doc.paragraphs[0].style='Title'
    for p in list(doc.paragraphs):
        if p.text.strip()=='PAGEBREAK':
            p.clear(); p.add_run().add_break(WD_BREAK.PAGE)
            p.paragraph_format.space_after=Pt(0); p.paragraph_format.space_before=Pt(0)
        elif p.text.startswith('EQUATIONNUMBER'):
            number=p.text.replace('EQUATIONNUMBER','').strip()
            previous=p._p.getprevious()
            if previous is not None:
                from docx.text.paragraph import Paragraph
                eq=Paragraph(previous,p._parent)
                eq.paragraph_format.keep_together=True
                eq.paragraph_format.tab_stops.add_tab_stop(Inches(6.4),WD_TAB_ALIGNMENT.RIGHT)
                eq.add_run('\t('+number+')').font.size=Pt(10)
                p._p.getparent().remove(p._p)
        elif p.style.name.startswith('Heading'):
            p.paragraph_format.keep_with_next=True
        if p._p.xpath('.//m:oMathPara'):
            p.paragraph_format.keep_together=True
        for run in p.runs:
            run.font.name='Times New Roman'
            if run.font.color.type is not None:
                run.font.color.rgb=RGBColor(0,0,0)
    style_tables(doc)
    # No automatic page numbering or title borders.
    for p in doc.paragraphs:
        for border in p._p.xpath('./w:pPr/w:pBdr'):
            border.getparent().remove(border)
    doc.save(target)
    return target

figure_map=[('Figure_1_regional_prevalence','figures/figure2_regional_prevalence.svg'),('Figure_S1_concentration_curve','figures/figure1_concentration_curve.svg'),('Figure_S2_wealth_gradient','figures/figure3_wealth_gradient.svg'),('Figure_S3_decomposition','figures/figure4_decomposition.svg')]
for stem,source in figure_map:
    png=FIGS/(stem+'.png')
    subprocess.run(['rsvg-convert','-w','3900','-o',str(png),str(ROOT/source)],check=True)
    with Image.open(png) as img:
        image=Image.new('RGB',img.size,'white')
        if img.mode=='RGBA': image.paste(img,mask=img.getchannel('A'))
        else: image.paste(img.convert('RGB'))
        image.save(png,dpi=(600,600))
        image.save(FIGS/(stem+'.tiff'),compression='tiff_lzw',dpi=(600,600))

main=read('submission/jph_main_manuscript.md')
abstract=main.split('\n## Abstract\n')[1].split('\nKeywords:')[0].replace('### ','## ').strip()
keywords='Keywords:'+main.split('\nKeywords:')[1].split('\n')[0]
body='## Introduction\n'+main.split('\n## Introduction\n')[1].split('\n## Main displays\n')[0]
body=body.replace('## References\n','PAGEBREAK\n\n## References\n')
body=re.sub(r'(?m)^(\d+)\. ',r'\n\1\\. ',body)
body=body.replace('All analytic decisions, numerical outputs, interpretations and final manuscript text were reviewed and verified by the authors, who take full responsibility for the work.','All analytic decisions, numerical outputs, interpretations and final manuscript text were reviewed and verified by the authors, who take full responsibility for the work.')
front='# '+TITLE+'\n\n'+NAMES+'\n\n'+AFFILIATION+'\n\nCorresponding author: Gideon Ofosu Kyere\n\nEmail: kyereofosu2003@gmail.com\n\n'+keywords
main_md=front+'\n\nPAGEBREAK\n\n# Abstract\n\n'+abstract+'\n\nPAGEBREAK\n\n'+body
main_md+='\n\nPAGEBREAK\n\n# Table I\n\nParticipant characteristics and survey-weighted anaemia prevalence\n\n'+block_table(read('results/table1_weighted_characteristics.md'))+'\nAnalytic sample: 7,557 women; BMI categories use 7,550 women because seven BMI values were unavailable or flagged. Percentages and prevalence incorporate sampling weights, strata and primary sampling units. CI denotes confidence interval.\n'
main_md+='\nPAGEBREAK\n\n# Table II\n\nSurvey-weighted adjusted associations with anaemia\n\n'+block_table(read('results/table2_full_adjusted_model.md'))+'\nAnalytic sample: 7,550 women. The model includes all listed factors. Adjusted odds ratios and confidence intervals use the stratified PSU sandwich covariance. Overall factor tests are in Supplementary Table S1.\n'
table3=read('results/table3_inequality_and_heterogeneity.md').split('**Aggregate sources:**')[0]
table3=table3.replace('# Table III. Wealth-related inequality decomposition and community-level heterogeneity','# Table III\n\nWealth-related inequality decomposition and community-level heterogeneity')
table3=table3.replace('## Panel A.','## Panel A').replace('## Panel B.','## Panel B')
table3=re.sub(r'The decomposition uses[\s\S]*?(?=## Panel B)', 'Panel A: 7,550 BMI-complete women; observed Erreygers index = -0.058853. Positive absolute contributions offset the negative index.\n\n',table3)
table3=re.sub(r'The null model includes[\s\S]*$', '**Notes:** Panel A uses a survey-weighted linear probability model with percentile intervals from 200 stratified PSU bootstrap replicates. Contributions are statistical, not causal; BMI does not identify anaemia aetiology. Panel B uses unweighted variational Bayes. PCV is descriptive because model samples differ. Approximate uncertainty and diagnostics are in the supplement. ICC = intraclass correlation coefficient; SD = standard deviation; PCV = proportional change in variance. Neither panel estimates an intervention effect.\n',table3)
main_md+='\nPAGEBREAK\n\n'+table3
legend=main.split('**Figure 1. Survey-weighted anaemia prevalence by region.**\n')[1].split('\nSource:')[0]
main_md+='\nPAGEBREAK\n\n# Figure legend\n\n## Figure 1\n\nSurvey-weighted anaemia prevalence by region\n\n'+legend
main_md+='\nPAGEBREAK\n\n# Figure 1\n\n!['+'Regional anaemia prevalence with 95% confidence intervals'+']('+str(FIGS/'Figure_1_regional_prevalence.png')+'){width=6.25in}\n'
make_doc('Main_Manuscript',main_md)

title=read('submission/jph_title_page.md').replace('# Title Page','# '+TITLE).replace('## Title\n\n**'+TITLE+'**\n\n','')
make_doc('Title_Page',title)

supp=read('supplementary/supplementary_material.md')
supp=supp.replace('\\[','$$').replace('\\]','$$')
supp=re.sub(r'(?m)^Repository file:.*\n','',supp)
supp=re.sub(r'(?m)^Source:.*\n','',supp)
supp=supp.replace('See `docs/real_data_validation_2026-10-09.md` and `docs/multilevel_reproducibility_validation.md`.','The converged variational-Bayes model is the reported estimator; the failed MAP/Laplace fit is diagnostic only.')
# Keep the tables and prose together, then append full-page figure displays.
supp=supp.replace('## Supplementary material','# Supplementary material')
for index,(stem,source) in enumerate(figure_map[1:],1):
    supp+='\nPAGEBREAK\n\n# Supplementary Figure S'+str(index)+'\n\n![Supplementary Figure S'+str(index)+']('+str(FIGS/(stem+'.png'))+'){width=6.25in}\n'
make_doc('Supplementary_Material',supp)

items=[
('1a','Design','Abstract; Methods, Study design and data source','Cross-sectional design stated.'),
('1b','Balanced abstract','Abstract','149 words; design, main estimates and interpretation included.'),
('2','Rationale','Introduction','Prior determinants, spatial and temporal decomposition studies distinguished.'),
('3','Objectives','Introduction, final paragraph','Prevalence, inequality, decomposition, associations and clustering specified.'),
('4','Study design','Methods, Study design and data source','Secondary cross-sectional analysis.'),
('5','Setting and dates','Methods, Study design and data source; reference 3','Ghana and exact fieldwork dates reported; dates cross-checked against the cited 2022 GDHS analysis (reference 4).'),
('6a','Eligibility and selection','Methods, Study design and data source; Outcome','Women aged 15–49 with valid classification; survey cluster selection described.'),
('7','Variables','Methods, Outcome; Socioeconomic and explanatory variables','Outcome, wealth and prespecified covariates defined.'),
('8','Measurement','Methods; reference 3','DHS classification and covariate source identified. See survey report for field measurement protocol.'),
('9','Bias','Discussion, Limitations of this study','Selection, measurement, BMI proxy and weighting limitations stated.'),
('10','Study size','Methods; Results, Study population and prevalence','All eligible observations used; 7,557 primary and 7,550 BMI-complete.'),
('11','Quantitative variables','Methods; equations 1–8','BMI boundaries, ranking and decomposition functional forms specified.'),
('12a','Statistical methods','Methods; equations 1–8','Survey estimators, confounder adjustment, bootstrap and multilevel methods stated.'),
('12b','Subgroups and interactions','Methods, Sensitivity and missing-data analyses; supplement','Exploratory subgroup checks described; no interaction estimates claimed.'),
('12c','Missing data','Methods; Results','Seven BMI exclusions stated. Measurement absence versus survey testing subsample should be distinguished if fuller disposition is available.'),
('12d','Sampling strategy','Methods, Complex-survey analysis','Weights, PSU, strata, Taylor variance and domain handling described.'),
('12e','Sensitivity analyses','Methods; supplementary BMI narrative','Exploratory checks separated from unexecuted planned outcomes.'),
('13a','Participant numbers','Methods; Results; Tables I–III','15,014 interviewed; 7,557 primary; 7,550 BMI-complete.'),
('13b','Reasons for exclusion','Methods, Study design and data source','BMI exclusions reported. Full haemoglobin-testing eligibility/nonresponse disposition is not documented in this manuscript.'),
('13c','Participant flow','Methods and Results','Counts provided in prose; no flow diagram. Do not infer unreported exclusion reasons.'),
('14a','Participant characteristics','Table I','Demographic, socioeconomic, reproductive and BMI categories.'),
('14b','Missing observations','Methods; Table I','BMI missingness reported; most prespecified covariates reported complete.'),
('15','Outcome summaries','Results; Table I; Figure 1','Weighted prevalence and 95% confidence intervals provided.'),
('16a','Estimates and precision','Results; Tables II–III','Adjusted odds ratios, design-based intervals and bootstrap decomposition intervals distinguished.'),
('16b','Category boundaries','Methods; Tables I–II','BMI cutoffs and covariate categories shown.'),
('16c','Absolute estimates','Results; Table I','Absolute prevalence reported; odds ratios not labelled as risks.'),
('17','Other analyses','Table III; Supplementary material','Decomposition, clustering and exploratory sensitivity documented.'),
('18','Key results','Discussion, Main finding of this study','Findings linked to study objectives.'),
('19','Limitations','Discussion, Limitations of this study','Cross-sectional, measurement, selection and model limitations retained.'),
('20','Interpretation','Discussion; Conclusion','Statistical contributions and conditional associations interpreted non-causally.'),
('21','Generalisability','Discussion, Limitations of this study','Population and geographical limits stated.'),
('22','Funding','Funding; reference 3','No external funding declared for this secondary study, pending author confirmation. Original survey support is documented in FR387.'),
]
strobe='# STROBE cross sectional reporting checklist\n\n'+TITLE+'\n\nThis checklist maps the cross-sectional STROBE items to the assembled manuscript. Section and table locators are used because the journal requests no manuscript page numbers. Entries identify reporting gaps rather than certifying completeness.\n\nSource: https://www.strobe-statement.org/checklists/\n\n| Item | Topic | Manuscript location | Assessment |\n|---|---|---|---|\n'+'\n'.join('| '+' | '.join(row)+' |' for row in items)
make_doc('STROBE_Checklist',strobe)

review='''# Scientific review and author confirmation

The methodology equations and promoted main results are consistent with the executable specification and validated aggregate outputs at the reported precision. The revised gap is a focused extension of the cited literature, not a claim to be the first study. This review does not rerun DHS microdata or establish causal effects.

## Scientific checks

The weighting and prevalence equations match the survey implementation. Wealth ranks use weighted midranks for ties. The binary correction satisfies E = 4 times the weighted outcome mean times C, or eight times the weighted covariance. The linear-model decomposition adds covariate contributions and a residual on the same complete-case sample. Logistic odds ratios and intervals match the design-based covariance implementation. ICC, MOR and PCV agree with the community-model code.

Main Table III reproduces the validated decomposition and multilevel values. The negative complete-case denominator is -0.058853, distinct from the full-sample E of -0.0589. All smaller-domain intervals include zero. The BMI percentage is a statistical contribution, not causal mediation or a recommendation to gain weight. The multilevel analysis is unweighted; its null and adjusted samples differ, so PCV is descriptive.

## Recorded author details

| Author | Affiliation | Position | Contact |
|---|---|---|---|
| Gideon Ofosu Kyere | Statistics and Actuarial Science, KNUST, Kumasi, Ghana | Graduate Research and Teaching Assistant, as recorded | kyereofosu2003@gmail.com; corresponding author |
| Wilhemina Adoma Pels | Same department and institution | To confirm | To confirm if required |
| Prince Apaah | Same department and institution | To confirm | papaah1@st.knust.edu.gh; ORCID 0009-0004-9551-4879 |
| Clement Acheampong | Same department and institution | To confirm | acheampongc479@gmail.com; ORCID 0009-0008-8436-4544 |

The corresponding author reports that the latest revision suggestions came from the coauthors. The repository records approval of the earlier manuscript on 12 October 2026. Neither record establishes final approval of this newly assembled package. Public ORCID ownership has not been independently verified.

## Confirm before merging or submitting

- Each author approves the exact assembled main manuscript, supplement, tables and figures.
- Confirm the author order, current positions, affiliation postal address/code and corresponding-author contact.
- Supply agreed CRediT roles. No roles have been inferred and no contributions section has been invented.
- Confirm originality, exclusive submission, funding, conflicts, ethics and DHS data-access declarations, and the accuracy of the AI-use statement.
- Resolve the STROBE entry on haemoglobin-testing disposition against the official report if available.

## Final approval record

| Author | Final package approval | Date |
|---|---|---|
| Gideon Ofosu Kyere | Pending | |
| Wilhemina Adoma Pels | Pending | |
| Prince Apaah | Pending | |
| Clement Acheampong | Pending | |

No merge or journal submission has been performed. This confirmation sheet stays outside the manuscript and is not a journal-upload declaration.
'''
make_doc('Author_Review_and_Confirmation',review)

letter='''# Cover letter

Dear Editors,

Please consider our manuscript, **Socioeconomic and Geographic Inequalities in Anaemia Among Women of Reproductive Age in Ghana**, as an Original Paper in the Journal of Public Health.

Using the 2022 Ghana Demographic and Health Survey, we examine wealth-related inequality in anaemia and its measured contributors, alongside geographic variation and residual community heterogeneity. The paper complements existing prevalence and determinants studies by distinguishing inequality across the wealth ranking from conditional associations.

The main text presents the validated decomposition and separate community-model summaries, with uncertainty and non-causal interpretation. The findings are relevant to public-health monitoring because national averages can conceal socioeconomic and geographic differences.

The study uses de-identified secondary survey data. Restricted DHS microdata are not redistributed. A data-availability statement directs researchers to The DHS Program's application process.

OpenAI ChatGPT assisted with code development, workflow documentation, drafting and language editing. The manuscript includes an AI-use disclosure for author review.

Thank you for considering the manuscript.

Sincerely,

Gideon Ofosu Kyere

Corresponding author

Department of Statistics and Actuarial Science

Kwame Nkrumah University of Science and Technology

Kumasi, Ghana

kyereofosu2003@gmail.com
'''
make_doc('Cover_Letter',letter)

readme='''# Anaemia manuscript review package

Main_Manuscript.docx contains the clean paper with native Word equations, editable Tables I–III and the main figure. Title_Page.docx, Supplementary_Material.docx, STROBE_Checklist.docx and Cover_Letter.docx are separate package components. Author_Review_and_Confirmation.docx records the scientific checks and outstanding author attestations.

The figures folder contains main Figure 1 and supplementary Figures S1–S3 in PNG and TIFF at 600 dpi. The source graphics were rasterised without changing estimates or geometry.

The package is for final author review. Coauthor suggestions are incorporated, but final approval of these assembled files, positions, contributions and submission declarations remains pending. No DHS microdata is included. No merge or submission has been performed.
'''
(FINAL/'README.md').write_text(readme,encoding='utf8')

# Assert semantic content survives native DOCX conversion.
for path in FINAL.glob('*.docx'):
    with ZipFile(path) as z:
        xml=z.read('word/document.xml').decode()
    assert 'EQUATIONNUMBER' not in xml and 'PAGEBREAK' not in xml
with ZipFile(FINAL/'Main_Manuscript.docx') as z:
    xml=z.read('word/document.xml').decode()
assert xml.count('<m:oMathPara>')==8, 'Expected eight editable display equations'
assert len(Document(FINAL/'Main_Manuscript.docx').tables)==4, 'Tables I, II and two panels of III expected'
print('Built six DOCX files and four paired figure exports')

